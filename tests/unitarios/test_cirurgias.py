from __future__ import annotations

import pandas as pd

from atualizador_paineis.core.config import load_panel_config
from atualizador_paineis.paineis.cirurgias.module import (
    RAW_COLUMNS,
    TREATED_COLUMNS,
    auxiliary_quantity_update,
    merge_period,
    transform,
)


def raw_frame(date: str, professional: str = "MÉDICO") -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Profissional Solicitante": professional,
                "Data de Solicitação": pd.Timestamp(date),
                "Paciente": "PACIENTE",
                "CPF": "000",
                "Telefone": None,
                "Celular": None,
                "Convenio": "CONVENIO",
                "Unidade": "UNIDADE",
            }
        ],
        columns=RAW_COLUMNS,
    )


def test_transform_adds_month_year_and_quantity() -> None:
    result = transform(raw_frame("2026-06-10 12:30:00"))

    assert result.columns.tolist() == TREATED_COLUMNS
    assert result.loc[0, "Mês"] == "JUNHO"
    assert result.loc[0, "Ano"] == 2026
    assert result.loc[0, "QTD"] == 1


def test_reprocessing_period_replaces_existing_rows() -> None:
    may = raw_frame("2026-05-10", "MAIO")
    old_june = raw_frame("2026-06-10", "JUNHO ANTIGO")
    existing_raw = pd.concat([may, old_june], ignore_index=True)
    existing_treated = transform(existing_raw)
    corrected_june = raw_frame("2026-06-15", "JUNHO CORRIGIDO")

    merged_raw, merged_treated, _, _, _ = merge_period(
        existing_raw,
        existing_treated,
        corrected_june,
        transform(corrected_june),
        pd.Period("2026-06", freq="M"),
    )

    assert merged_raw["Profissional Solicitante"].tolist() == [
        "MAIO",
        "JUNHO CORRIGIDO",
    ]
    assert len(merged_treated) == 2
    assert set(merged_treated["Mês"]) == {"MAIO", "JUNHO"}
    assert set(merged_treated["QTD"]) == {1}


def test_auxiliary_quantity_update_targets_processed_month(tmp_path) -> None:
    panel = tmp_path / "painel.xlsx"
    pd.DataFrame(
        {
            "Mês": ["JANEIRO", "JUNHO", "JULHO"],
            "QTD": [10, 20, 30],
            "Porcentagem": [0.1, 0.2, 0.3],
        }
    ).to_excel(panel, sheet_name="INPUTS", index=False, startrow=9)

    from openpyxl import load_workbook

    workbook = load_workbook(panel)
    worksheet = workbook["INPUTS"]
    from openpyxl.worksheet.table import Table, TableStyleInfo

    table = Table(displayName="Tabela7", ref="A10:C13")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    worksheet.add_table(table)
    workbook.save(panel)

    update = auxiliary_quantity_update(
        panel,
        load_panel_config("cirurgias"),
        pd.Period("2026-07", freq="M"),
        42,
    )

    assert update.sheet_name == "INPUTS"
    assert update.table_name == "Tabela7"
    assert update.row_index == 2
    assert update.column_name == "QTD"
    assert update.value == 42
